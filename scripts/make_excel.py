"""
Build Radiative_Transfer_Lab.xlsx: an interactive spreadsheet companion to the course.

Every sheet has yellow input cells; all results, tables and charts are Excel
formulas, so they update as soon as an input changes (Excel, LibreOffice,
Google Sheets). Fixed mathematical tables (Hopf function, E2, Chandrasekhar
H-function, exact scattering solutions) come from lte_solvers.py and are stored
on the 'Data' sheet, clearly labelled.

    python3 scripts/make_excel.py
"""
import os
import sys

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from scipy.special import expn

sys.path.insert(0, os.path.dirname(__file__))
import lte_solvers as rt  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "..", "Radiative_Transfer_Lab.xlsx")

F = "Arial"
f_title = Font(name=F, size=14, bold=True, color="1F3A8A")
f_bold = Font(name=F, size=10, bold=True)
f_norm = Font(name=F, size=10)
f_note = Font(name=F, size=9, italic=True, color="555555")
f_input = Font(name=F, size=10, bold=True, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_out = PatternFill("solid", fgColor="E8EEF8")
fill_head = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="B0B7C3")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()


# ------------------------------------------------------------------ helpers
def sheet(name, title, concept, equation):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = f_title
    ws["A2"] = concept
    ws["A2"].font = f_norm
    ws["A3"] = "Equation: " + equation
    ws["A3"].font = Font(name=F, size=10, italic=True, color="1F3A8A")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 3
    for col in "EFGHIJKL":
        ws.column_dimensions[col].width = 14
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, cell, text):
    ws[cell] = text
    ws[cell].font = f_bold
    ws[cell].fill = fill_head


def inp(ws, row, label, value, unit="", note=None):
    ws.cell(row=row, column=1, value=label).font = f_norm
    c = ws.cell(row=row, column=2, value=value)
    c.font, c.fill, c.border = f_input, fill_in, box
    ws.cell(row=row, column=3, value=unit).font = f_note
    if note:
        c.comment = Comment(note, "course")
    return f"$B${row}"


def out(ws, row, label, formula, unit="", fmt="0.000E+00"):
    ws.cell(row=row, column=1, value=label).font = f_norm
    c = ws.cell(row=row, column=2, value=formula)
    c.font, c.fill, c.border, c.number_format = f_norm, fill_out, box, fmt
    ws.cell(row=row, column=3, value=unit).font = f_note
    return f"$B${row}"


def table(ws, col0, row0, headers, rows, fmts=None):
    """rows: list of lists (values or formulas). Returns (first_row, last_row)."""
    for j, h in enumerate(headers):
        c = ws.cell(row=row0, column=col0 + j, value=h)
        c.font, c.fill, c.border = f_bold, fill_head, box
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(row=row0 + 1 + i, column=col0 + j, value=v)
            c.font = f_norm
            if fmts and fmts[j]:
                c.number_format = fmts[j]
    return row0 + 1, row0 + len(rows)


def chart(ws, anchor, title, xlab, ylab, xcol, ycols, r1, r2, names, xlog=False, ylog=False,
          xmin=None, xmax=None, ymin=None, ymax=None, w=17, h=8.5, yfmt=None):
    ch = ScatterChart()
    ch.title = title
    ch.style = 2
    ch.x_axis.title = xlab
    ch.y_axis.title = ylab
    ch.width, ch.height = w, h
    ch.legend.position = "b"
    xref = Reference(ws, min_col=xcol, min_row=r1, max_row=r2)
    colors = ["2A78D6", "EB6834", "1BAF7A", "6250D6", "E34948"]
    for k, (yc, nm) in enumerate(zip(ycols, names)):
        yref = Reference(ws, min_col=yc, min_row=r1, max_row=r2)
        s = Series(yref, xref, title=nm)
        s.marker.symbol = "none"
        s.smooth = False
        s.graphicalProperties.line.solidFill = colors[k % len(colors)]
        s.graphicalProperties.line.width = 22000
        if k > 0:
            s.graphicalProperties.line.dashStyle = "dash"
        ch.series.append(s)
    if xlog:
        ch.x_axis.scaling.logBase = 10
    if ylog:
        ch.y_axis.scaling.logBase = 10
    for ax, lo, hi in ((ch.x_axis, xmin, xmax), (ch.y_axis, ymin, ymax)):
        if lo is not None:
            ax.scaling.min = lo
        if hi is not None:
            ax.scaling.max = hi
        ax.delete = False
        ax.majorGridlines = ax.majorGridlines
    ch.x_axis.number_format = "General"
    ch.y_axis.number_format = yfmt or ("0E+00" if ylog else "General")
    ws.add_chart(ch, anchor)


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ------------------------------------------------------------------ Start here
ws0 = wb.active
ws0.title = "Start here"
ws0.sheet_view.showGridLines = False
ws0.column_dimensions["A"].width = 26
ws0.column_dimensions["B"].width = 95
ws0["A1"] = "Radiative Transfer Lab (spreadsheet edition)"
ws0["A1"].font = Font(name=F, size=16, bold=True, color="1F3A8A")
ws0["A2"] = "Companion to the two courses 'Radiative Transfer' (RT1-RT6) and 'Stellar Atmospheres' (SA1-SA5). One sheet per experiment, labelled with its lecture."
ws0["A2"].font = f_norm
ws0["A4"] = "How to use"
ws0["A4"].font = f_bold
legend = [
    ("Yellow cell, blue text", "INPUT: type a new value here. Everything else on the sheet updates automatically."),
    ("Light-blue cell", "RESULT: a formula computed from the inputs (do not type over it)."),
    ("Tables and charts", "Computed from the inputs with formulas; the charts redraw when you change an input."),
    ("Cell comments", "Hover over an input cell to see what it means and a sensible range."),
    ("Example values", "Each sheet opens with realistic example values (e.g. the Sun: T_eff = 5772 K)."),
]
for i, (a, b) in enumerate(legend):
    ws0.cell(row=5 + i, column=1, value=a).font = f_bold
    ws0.cell(row=5 + i, column=2, value=b).font = f_norm
ws0["A5"].fill = fill_in
ws0["A5"].font = f_input
ws0["A6"].fill = fill_out
ws0["A11"] = "Sheets"
ws0["A11"].font = f_bold
SHEETS = [
    ("RT1 Optical depth", "Beam attenuation, optical depth, magnitudes of extinction"),
    ("RT1 Intensity", "Flux falls as 1/d^2, solid angle too: intensity is conserved"),
    ("RT2 Planck", "Black-body spectrum, Wien peak, Stefan-Boltzmann, RJ and Wien limits"),
    ("RT2 Random walk", "Mean free path, photon diffusion time, thermalisation depth"),
    ("RT3 Slab & MIT star", "Slab solution; MIT two-layer star with Balmer + Paschen lines; thermal inversion"),
    ("SA1 Saha-Boltzmann", "Hydrogen ionisation and n=2 population; why Balmer lines peak in A stars"),
    ("RT4 Eddington-Barbier", "Emergent intensity, limb darkening, Eddington-Barbier error"),
    ("RT4 Scattering", "S = eps B + (1-eps) J: Eddington solution, sqrt(eps) law; Rosseland mean calculator"),
    ("SA4 Gray atmosphere", "Exact gray LTE model: T(tau), limb darkening at any wavelength, emergent spectrum, pressure"),
    ("SA4 Lines", "Milne-Eddington LTE line profile and curve of growth"),
    ("Constants", "Physical constants used by all formulas (named cells)"),
    ("Data", "Fixed mathematical tables from scripts/lte_solvers.py (Hopf q, E2, H-function, exact scattering)"),
]
for i, (a, b) in enumerate(SHEETS):
    ws0.cell(row=12 + i, column=1, value=a).font = f_link
    ws0.cell(row=12 + i, column=2, value=b).font = f_norm
    ws0.cell(row=12 + i, column=1).hyperlink = f"#'{a}'!A1"
ws0["A25"] = "An interactive web version with sliders (same physics) is in app/index.html of the course repository."
ws0["A25"].font = f_note

# ------------------------------------------------------------------ Constants
wc = wb.create_sheet("Constants")
wc.sheet_view.showGridLines = False
wc.column_dimensions["A"].width = 34
wc.column_dimensions["B"].width = 18
wc.column_dimensions["C"].width = 16
wc.column_dimensions["D"].width = 12
wc["A1"] = "Physical constants (SI, CODATA 2018)"
wc["A1"].font = f_title
CONST = [("Planck constant h", "hP", 6.62607015e-34, "J s"), ("speed of light c", "cL", 2.99792458e8, "m/s"),
         ("Boltzmann constant k", "kB", 1.380649e-23, "J/K"), ("Stefan-Boltzmann sigma", "sigSB", 5.670374e-8, "W m^-2 K^-4"),
         ("electron mass m_e", "mE", 9.1093837e-31, "kg"), ("electron volt", "eV", 1.602176634e-19, "J"),
         ("hydrogen mass m_H", "mH", 1.6735e-27, "kg"), ("atomic mass unit", "amu", 1.66053907e-27, "kg"),
         ("classical electron radius r_e", "rE", 2.8179403e-15, "m"), ("Thomson cross-section", "sigT", 6.6524587e-29, "m^2"),
         ("solar radius", "Rsun", 6.957e8, "m"), ("solar luminosity", "Lsun", 3.828e26, "W"),
         ("astronomical unit", "AU", 1.495978707e11, "m"), ("parsec", "pc", 3.0857e16, "m"), ("year", "yr", 3.15576e7, "s")]
table(wc, 1, 3, ["Constant", "Value", "Unit", "Name"], [[a, v, u, n] for a, n, v, u in CONST], [None, "0.000000E+00", None, None])
for i, (a, n, v, u) in enumerate(CONST):
    name(n, f"Constants!$B${4 + i}")

# ------------------------------------------------------------------ Data sheet
wd = wb.create_sheet("Data")
wd.sheet_view.showGridLines = False
wd["A1"] = "Fixed mathematical tables (not inputs). Source: scripts/lte_solvers.py of the course repository."
wd["A1"].font = f_bold
t, q = rt.hopf_function()
TAU = np.concatenate(([0.0], np.logspace(-3, np.log10(30.0), 70)))
QQ = np.interp(TAU, t, q)
E2 = np.array([1.0] + [float(expn(2, x)) for x in TAU[1:]])
hop_rows = [[float(a), float(b), float(c)] for a, b, c in zip(TAU, QQ, E2)]
r1h, r2h = table(wd, 1, 3, ["tau", "Hopf q(tau) (exact)", "E2(tau)"], hop_rows, ["0.0000E+00", "0.000000", "0.000000E+00"])
MU = np.linspace(0, 1, 21)
Hm = rt.chandrasekhar_H(MU)
r1m, r2m = table(wd, 5, 3, ["mu", "H(mu) exact", "I(mu)/I(1) exact gray"], [[float(m), float(h), float(h / Hm[-1])] for m, h in zip(MU, Hm)],
                 ["0.00", "0.0000", "0.0000"])
eps_list = [1.0, 0.3, 0.1, 0.03, 0.01, 0.003, 0.001]
ex_rows = []
for e in eps_list:
    if e == 1.0:
        ex_rows.append([e, 1.0, 1.0])
        continue
    tt, S = rt.solve_scattering_atmosphere(e)
    ex_rows.append([e, float(S[0]), float(np.sqrt(e))])
table(wd, 9, 3, ["eps", "exact S(0)/B (numerical)", "sqrt(eps)"], ex_rows, ["0.000", "0.0000", "0.0000"])
for col, wdt in zip("ABCDEFGHIJK", [10, 18, 14, 3, 8, 12, 18, 3, 8, 22, 10]):
    wd.column_dimensions[col].width = wdt
TAU_R = f"Data!$A${r1h}:$A${r2h}"
Q_R = f"Data!$B${r1h}:$B${r2h}"
E2_R = f"Data!$C${r1h}:$C${r2h}"
NH = len(TAU)

# ------------------------------------------------------------------ L1
ws = sheet("RT1 Optical depth", "Lecture RT1 - How much light gets through a slab of gas?",
           "Absorption removes a fixed fraction of the beam per unit length. Optical depth counts mean free paths.",
           "I(s) = I(0) exp(-tau),  tau = alpha * s,  A = 1.086 tau  [mag]")
header(ws, "A5", "Inputs")
a = inp(ws, 6, "Absorption coefficient alpha", 0.5, "per km", "Fraction of the beam removed per km. Try 0.01 to 10.")
L = inp(ws, 7, "Path length L", 4, "km", "Total path through the gas. Try 0.1 to 20.")
header(ws, "A9", "Results")
tau = out(ws, 10, "Optical depth tau = alpha L", f"={a}*{L}", "", "0.000")
out(ws, 11, "Transmitted fraction exp(-tau)", f"=EXP(-{tau})", "", "0.0000")
out(ws, 12, "Extinction A = 2.5 log10(e^tau)", f"=1.0857*{tau}", "mag", "0.000")
out(ws, 13, "Mean free path 1/alpha", f"=1/{a}", "km", "0.000")
out(ws, 14, "Regime", f'=IF({tau}<0.3,"optically thin",IF({tau}<3,"intermediate (tau ~ 1)","optically thick"))', "", "@")
rows = [[f"={L}*1.5*{i}/40", f"=EXP(-{a}*E{17 + i})"] for i in range(41)]
r1, r2 = table(ws, 5, 16, ["s [km]", "I(s)/I(0)"], rows, ["0.00", "0.0000"])
chart(ws, "H5", "Fraction of the beam left after distance s", "distance s [km]", "I(s)/I(0)", 5, [6], r1, r2, ["I/I0"], ymin=0, ymax=1)

# ------------------------------------------------------------------ L2
ws = sheet("RT1 Intensity", "Lecture RT1 - Why a star looks equally bright per square arcsecond from any distance",
           "Flux falls as 1/d^2, but the solid angle of the star falls in exactly the same way: their ratio, the intensity, is constant.",
           "F(d) = sigma T^4 (R/d)^2,  Omega = pi (R/d)^2,  I = F/Omega = sigma T^4 / pi")
header(ws, "A5", "Inputs")
T = inp(ws, 6, "Effective temperature T_eff", 5772, "K", "Sun 5772 K; Sirius A 9940 K.")
R = inp(ws, 7, "Stellar radius R", 1, "R_sun", "Sun 1; Sirius A 1.711.")
d = inp(ws, 8, "Distance d", 1, "AU", "Earth 1 AU, Mars 1.524 AU, Sirius 543900 AU.")
header(ws, "A10", "Results")
Fs = out(ws, 11, "Surface flux sigma T^4", f"=sigSB*{T}^4", "W m^-2")
rat = out(ws, 12, "(R/d)^2", f"=({R}*Rsun/({d}*AU))^2", "")
out(ws, 13, "Flux received F", f"={Fs}*{rat}", "W m^-2")
out(ws, 14, "Solid angle Omega", f"=PI()*{rat}", "sr")
out(ws, 15, "Intensity I = F / Omega", f"=B13/B14", "W m^-2 sr^-1")
out(ws, 16, "Same as sigma T^4 / pi ?", f"={Fs}/PI()", "W m^-2 sr^-1")
out(ws, 17, "Angular diameter", f"=2*ASIN(SQRT({rat}))*206264.806", "arcsec", "0.000")
out(ws, 18, "Luminosity", f"=4*PI()*({R}*Rsun)^2*{Fs}/Lsun", "L_sun", "0.000")
dvals = np.logspace(-2, 6, 33)
rows = [[float(x), f"=({R}*Rsun/(E{21 + i}*AU))^2", f"=F{21 + i}", "=1"] for i, x in enumerate(dvals)]
r1, r2 = table(ws, 5, 20, ["d [AU]", "F / F_surface", "Omega / pi", "I / I_surface"], rows, ["0.00E+00", "0.00E+00", "0.00E+00", "0"])
chart(ws, "J5", "Relative to the stellar surface (log-log)", "distance d [AU]", "value / surface value", 5, [6, 8], r1, r2,
      ["flux F and solid angle (identical)", "intensity I"], xlog=True, ylog=True, ymax=100)

# ------------------------------------------------------------------ L3
ws = sheet("RT2 Planck", "Lecture RT2 - How a black body changes with temperature",
           "Raise T: the peak moves to shorter wavelengths (Wien) and the total flux grows as T^4 (Stefan-Boltzmann).",
           "B_lambda(T) = 2hc^2/lambda^5 / (exp(hc/(lambda k T)) - 1);  lambda_max T = 2.898e-3 m K;  F = sigma T^4")
header(ws, "A5", "Inputs")
T = inp(ws, 6, "Temperature T", 5772, "K", "Try 3000 (red dwarf), 5772 (Sun), 10000 (A star), 30000 (O star).")
T2 = inp(ws, 7, "Comparison temperature", 5772, "K", "Second curve on the chart.")
header(ws, "A9", "Results")
out(ws, 10, "Peak wavelength lambda_max", f"=2.897771955E-3/{T}*1E9", "nm", "0.0")
out(ws, 11, "Peak frequency nu_max = 5.879e10 T", f"=5.879E10*{T}", "Hz")
out(ws, 12, "Surface flux sigma T^4", f"=sigSB*{T}^4", "W m^-2")
out(ws, 13, "Flux relative to the Sun", f"=({T}/5772)^4", "", "0.000")
out(ws, 14, "B_lambda at 500 nm", f"=2*hP*cL^2/(500E-9)^5/(EXP(hP*cL/(500E-9*kB*{T}))-1)", "W m^-3 sr^-1")
out(ws, 15, "x = hc/(lambda k T) at 500 nm", f"=hP*cL/(500E-9*kB*{T})", "", "0.000")
lam = np.round(np.logspace(np.log10(100), np.log10(5000), 60), 2)
rows = []
for i, l in enumerate(lam):
    r = 18 + i
    rows.append([float(l), f"=2*hP*cL^2/(E{r}*1E-9)^5/(EXP(hP*cL/(E{r}*1E-9*kB*{T}))-1)",
                 f"=2*hP*cL^2/(E{r}*1E-9)^5/(EXP(hP*cL/(E{r}*1E-9*kB*{T2}))-1)",
                 f"=2*cL*kB*{T}/(E{r}*1E-9)^4", f"=2*hP*cL^2/(E{r}*1E-9)^5*EXP(-hP*cL/(E{r}*1E-9*kB*{T}))"])
r1, r2 = table(ws, 5, 17, ["lambda [nm]", "B_lambda(T)", "B_lambda(T2)", "Rayleigh-Jeans", "Wien"], rows,
               ["0.0", "0.000E+00", "0.000E+00", "0.000E+00", "0.000E+00"])
chart(ws, "K5", "Planck function B_lambda (linear axes)", "wavelength [nm]", "B_lambda [W m^-3 sr^-1]", 5, [6, 7], r1, r2,
      ["B_lambda(T)", "B_lambda(comparison T)"], xmin=100, xmax=3000, yfmt="0.0E+00")
chart(ws, "K23", "Planck vs Rayleigh-Jeans and Wien limits (log-log)", "wavelength [nm]", "B_lambda", 5, [6, 8, 9], r1, r2,
      ["Planck", "Rayleigh-Jeans", "Wien"], xlog=True, ylog=True)

# ------------------------------------------------------------------ L4
ws = sheet("RT2 Random walk", "Lecture RT2 - How long does light take to leave a star?",
           "A scattered photon random-walks: after N steps of length l it is only sqrt(N) l away, so escaping from optical depth tau takes N ~ tau^2 steps.",
           "l = 1/(n sigma),  tau = R/l,  N = tau^2,  t = R^2/(l c),  thermalisation depth ~ 1/sqrt(eps)")
header(ws, "A5", "Inputs (default: uniform Sun, Thomson scattering)")
M = inp(ws, 6, "Mass", 1, "M_sun", "Mass of the sphere in solar masses.")
Rr = inp(ws, 7, "Radius", 1, "R_sun")
X = inp(ws, 8, "Hydrogen mass fraction X", 0.7, "", "Sets the electrons per unit mass: mu_e = 2/(1+X).")
eps = inp(ws, 9, "Destruction probability eps", 0.001, "", "Probability that an interaction destroys the photon (absorption).")
header(ws, "A11", "Results")
rho = out(ws, 12, "Mean density", f"={M}*1.989E30/(4/3*PI()*({Rr}*Rsun)^3)", "kg m^-3", "0.0")
ne = out(ws, 13, "Electron density n_e = rho/(mu_e m_H)", f"={rho}/((2/(1+{X}))*mH)", "m^-3")
ell = out(ws, 14, "Mean free path l = 1/(n_e sigma_T)", f"=1/({ne}*sigT)", "m", "0.0000")
tt = out(ws, 15, "Optical radius tau = R/l", f"={Rr}*Rsun/{ell}", "")
out(ws, 16, "Number of scatterings N = tau^2", f"={tt}^2", "")
out(ws, 17, "Diffusion time R^2/(l c)", f"=({Rr}*Rsun)^2/({ell}*cL)/yr", "years", "#,##0")
out(ws, 18, "Light-crossing time R/c", f"={Rr}*Rsun/cL", "s", "0.00")
out(ws, 19, "Thermalisation depth 1/sqrt(eps)", f"=1/SQRT({eps})", "", "0.0")
out(ws, 20, "Scatterings before destruction 1/eps", f"=1/{eps}", "", "#,##0")
rows = [[float(x), f"=E{23 + i}^2", f"=MIN(1,EXP(-E{23 + i}*SQRT(3*{eps})))"] for i, x in enumerate(np.arange(1, 31))]
r1, r2 = table(ws, 5, 22, ["cloud tau", "N = tau^2", "escape chance ~ exp(-tau sqrt(3 eps))"], rows, ["0", "#,##0", "0.0000"])
chart(ws, "I5", "Scatterings needed to escape: N = tau^2", "optical radius tau", "N", 5, [6], r1, r2, ["N = tau^2"])
chart(ws, "I23", "Rough escape probability with absorption (diffusion estimate)", "optical radius tau", "probability", 5, [7], r1, r2,
      ["exp(-tau sqrt(3 eps))"], ymin=0, ymax=1)

# ------------------------------------------------------------------ L5
ws = sheet("RT3 Slab & MIT star", "Lecture RT3 - Absorption or emission lines? A slab and the MIT 8.901 two-layer star",
           "The beam relaxes toward the source function. Lines are dark if the gas in front is cooler than the background, bright if hotter.",
           "I = I0 exp(-tau) + S (1 - exp(-tau));  I = B(T_base) exp(-tau) + B(T_top) (1 - exp(-tau))")
header(ws, "A5", "Part A - uniform slab")
I0 = inp(ws, 6, "Incoming intensity I(0)/S", 2.2, "", "Brightness of the beam entering the slab, in units of S.")
ts = inp(ws, 7, "Slab optical depth tau", 1.5, "")
out(ws, 8, "Emergent I/S", f"={I0}*EXP(-{ts})+(1-EXP(-{ts}))", "", "0.0000")
header(ws, "A10", "Part B - MIT two-layer star")
Tb = inp(ws, 11, "T_base (thick interior)", 6500, "K", "Hot, optically thick base.")
Tt = inp(ws, 12, "T_top (thin upper layer)", 4500, "K", "Make it larger than T_base for a thermal inversion.")
t0 = inp(ws, 13, "Line strength tau0 (H-alpha)", 2, "", "Optical depth at the centre of H-alpha. Try 0.1 to 100.")
tc = inp(ws, 14, "Continuum optical depth tau_c", 0.02, "")
bg = inp(ws, 15, "Background: 1 = thick base, 0 = dark sky", 1, "", "Use 0 for a thin cloud seen against empty sky.")
header(ws, "A17", "Results at H-alpha (656.47 nm)")
Bb = f"(2*hP*cL^2/(656.47E-9)^5/(EXP(hP*cL/(656.47E-9*kB*{Tb}))-1))"
Btp = f"(2*hP*cL^2/(656.47E-9)^5/(EXP(hP*cL/(656.47E-9*kB*{Tt}))-1))"
# Line list (Rydberg) on this sheet
lines = []
for n1 in (2, 3):
    for n2 in range(n1 + 1, n1 + 12):
        lam0 = 1e9 / (1.09678e7 * (1 / n1**2 - 1 / n2**2))
        rel = ((n1 + 1) / n2) ** 3 * (1.0 if n1 == 2 else 0.6)
        lines.append([n1, n2, round(lam0, 3), round(rel, 5)])
lr1, lr2 = table(ws, 13, 5, ["n1", "n2", "lambda0 [nm] (Rydberg)", "relative strength"], lines, ["0", "0", "0.00", "0.0000"])
for col, wdt in zip("MNOP", [6, 6, 16, 14]):
    ws.column_dimensions[col].width = wdt
LAM_R, REL_R = f"$O${lr1}:$O${lr2}", f"$P${lr1}:$P${lr2}"
tl = out(ws, 18, "tau at line centre", f"={tc}+{t0}*SUMPRODUCT({REL_R},EXP(0-((656.47-{LAM_R})/1.2)^2))", "", "0.000")
Ic = out(ws, 19, "Continuum intensity", f"={bg}*{Bb}*EXP(-{tc})+{Btp}*(1-EXP(-{tc}))", "W m^-3 sr^-1")
Il = out(ws, 20, "Line-centre intensity", f"={bg}*{Bb}*EXP(-{tl})+{Btp}*(1-EXP(-{tl}))", "W m^-3 sr^-1")
out(ws, 21, "Line depth 1 - I_line/I_cont", f"=1-{Il}/{Ic}", "(negative = emission)", "0.000")
out(ws, 22, "Line type", f'=IF({Il}<{Ic},"absorption line","emission line")', "", "@")
rows = []
for i in range(31):
    r = 26 + i
    rows.append([f"=6*{i}/30", f"={I0}*EXP(-E{r})+(1-EXP(-E{r}))", "=1"])
r1a, r2a = table(ws, 5, 25, ["tau", "I/S", "S"], rows, ["0.00", "0.0000", "0"])
chart(ws, "R5", "Part A: intensity relaxes toward S", "optical depth into slab", "I / S", 5, [6, 7], r1a, r2a, ["I/S", "S"], ymin=0)
rows = []
lam = np.arange(360, 1001, 2.0)
for i, l in enumerate(lam):
    r = 26 + i
    Bbr = f"(2*hP*cL^2/(H{r}*1E-9)^5/(EXP(hP*cL/(H{r}*1E-9*kB*{Tb}))-1))"
    Btr = f"(2*hP*cL^2/(H{r}*1E-9)^5/(EXP(hP*cL/(H{r}*1E-9*kB*{Tt}))-1))"
    rows.append([float(l), f"={tc}+{t0}*SUMPRODUCT({REL_R},EXP(0-((H{r}-{LAM_R})/1.2)^2))",
                 f"=({bg}*{Bbr}*EXP(-I{r})+{Btr}*(1-EXP(-I{r})))/1E12", f"={bg}*{Bbr}/1E12", f"={Btr}/1E12"])
r1b, r2b = table(ws, 8, 25, ["lambda [nm]", "tau(lambda)", "I [1e12]", "B(T_base) x bg [1e12]", "B(T_top) [1e12]"], rows,
                 ["0", "0.000", "0.00", "0.00", "0.00"])
chart(ws, "R24", "Part B: emergent spectrum (Balmer + Paschen lines, widths exaggerated)", "wavelength [nm]", "I [1e12 W m^-3 sr^-1]",
      8, [10, 11, 12], r1b, r2b, ["emergent I", "B(T_base)", "B(T_top)"], xmin=360, xmax=1000, w=26, h=10)

# ------------------------------------------------------------------ L6
ws = sheet("SA1 Saha-Boltzmann", "Lecture SA1 - Why are Balmer lines strongest in A stars?",
           "In LTE, excitation follows Boltzmann and ionisation follows Saha. The n=2 population (Balmer absorbers) peaks near 10 000 K.",
           "N_II n_e / N_I = (2 U_II/U_I)(2 pi m_e k T/h^2)^1.5 exp(-13.6 eV/kT);  N2/N1 = 4 exp(-10.2 eV/kT)")
header(ws, "A5", "Inputs")
T = inp(ws, 6, "Temperature T", 9900, "K", "Try 5772 (Sun), 9900 (A0 star), 20000 (B star).")
Pe = inp(ws, 7, "Electron pressure P_e", 20, "Pa", "Typical photospheres: 1 to 100 Pa.")
header(ws, "A9", "Results")
neL6 = out(ws, 10, "Electron density n_e = P_e/kT", f"={Pe}/(kB*{T})", "m^-3")
out(ws, 11, "kT", f"=kB*{T}/eV", "eV", "0.0000")
r21 = out(ws, 12, "N2/N1 (Boltzmann)", f"=4*EXP(-10.199*eV/(kB*{T}))", "")
xs = out(ws, 13, "N_II/N_I (Saha)", f"=(2*PI()*mE*kB*{T}/hP^2)^1.5*EXP(-13.598*eV/(kB*{T}))/{neL6}", "")
out(ws, 14, "Ionised fraction", f"={xs}/(1+{xs})", "", "0.0000")
out(ws, 15, "N2 / N_total (Balmer absorbers)", f"={r21}/(1+{r21})/(1+{xs})", "")
rows = []
Ts = np.arange(3000, 30001, 100)
for i, tv in enumerate(Ts):
    r = 18 + i
    sa = f"(2*PI()*mE*kB*E{r}/hP^2)^1.5*EXP(-13.598*eV/(kB*E{r}))/({Pe}/(kB*E{r}))"
    rows.append([int(tv), f"={sa}/(1+{sa})", f"=1E5*(4*EXP(-10.199*eV/(kB*E{r})))/(1+4*EXP(-10.199*eV/(kB*E{r})))/(1+{sa})"])
r1, r2 = table(ws, 5, 17, ["T [K]", "ionised fraction", "N2/N_total x 1e5"], rows, ["0", "0.0000", "0.0000"])
out(ws, 16, "Temperature of the Balmer peak", f"=INDEX(E{r1}:E{r2},MATCH(MAX(G{r1}:G{r2}),G{r1}:G{r2},0))", "K", "0")
chart(ws, "I5", "Ionised fraction of hydrogen", "temperature [K]", "N_II / N_total", 5, [6], r1, r2, ["ionised fraction"], ymin=0, ymax=1)
chart(ws, "I23", "Fraction in n = 2 (x 1e5): Balmer line strength", "temperature [K]", "N2/N_total x 1e5", 5, [7], r1, r2, ["N2/N_total x 1e5"], ymin=0)

# ------------------------------------------------------------------ L7
ws = sheet("RT4 Eddington-Barbier", "Lecture RT4 - What depth do we see at each point of the stellar disk?",
           "The emergent intensity is S averaged with weight exp(-tau/mu)/mu. Eddington-Barbier: I(0,mu) ~ S(tau = mu), exact for linear S.",
           "S(tau) = a + b tau + c tau^2  ->  I(0,mu) = a + b mu + 2 c mu^2;  F = pi (a + 2b/3 + c) ~ pi S(2/3)")
header(ws, "A5", "Inputs")
a = inp(ws, 6, "a (surface value of S)", 0.4, "", "Solar 500 nm: a = 0.36 with b = 0.64.")
b = inp(ws, 7, "b (gradient)", 0.6, "", "b > 0: S increases inward -> limb darkening. b < 0: limb brightening.")
c = inp(ws, 8, "c (curvature)", 0, "", "c = 0 makes Eddington-Barbier exact.")
mu = inp(ws, 9, "mu = cos(theta) (1 = centre, 0 = limb)", 0.5, "")
header(ws, "A11", "Results")
Iex = out(ws, 12, "Exact I(0,mu) = a + b mu + 2 c mu^2", f"={a}+{b}*{mu}+2*{c}*{mu}^2", "", "0.0000")
Ieb = out(ws, 13, "Eddington-Barbier S(tau = mu)", f"={a}+{b}*{mu}+{c}*{mu}^2", "", "0.0000")
out(ws, 14, "Error of Eddington-Barbier", f"=IF({Iex}=0,0,({Ieb}-{Iex})/{Iex})", "", "0.00%")
out(ws, 15, "Flux pi (a + 2b/3 + c)", f"=PI()*({a}+2*{b}/3+{c})", "", "0.0000")
out(ws, 16, "pi S(2/3)", f"=PI()*({a}+2*{b}/3+{c}*4/9)", "", "0.0000")
out(ws, 17, "Limb/centre ratio I(0)/I(1)", f"={a}/({a}+{b}+2*{c})", "", "0.000")
rows = [[f"={i}/20", f"=({a}+{b}*E{20 + i}+2*{c}*E{20 + i}^2)/({a}+{b}+2*{c})"] for i in range(21)]
r1, r2 = table(ws, 5, 19, ["mu", "I(mu)/I(1)"], rows, ["0.00", "0.0000"])
rows = [[f"=3*{i}/40", f"={a}+{b}*H{20 + i}+{c}*H{20 + i}^2", f"=EXP(-H{20 + i}/{mu})/{mu}"] for i in range(41)]
r1b, r2b = table(ws, 8, 19, ["tau", "S(tau)", "weight exp(-tau/mu)/mu"], rows, ["0.000", "0.0000", "0.0000"])
chart(ws, "L5", "Limb darkening I(mu)/I(1)", "mu = cos theta", "I(mu)/I(1)", 5, [6], r1, r2, ["I(mu)/I(1)"], ymin=0, xmin=0, xmax=1)
chart(ws, "L23", "Source function S(tau)", "vertical optical depth tau", "S", 8, [9], r1b, r2b, ["S(tau)"])
chart(ws, "L41", "Weight: where the light at this mu comes from", "vertical optical depth tau", "weight", 8, [10], r1b, r2b, ["exp(-tau/mu)/mu"])

# ------------------------------------------------------------------ L8
ws = sheet("RT4 Scattering", "Lecture RT4 - Why is S below B at the surface of an isothermal scattering atmosphere?",
           "Photons scatter ~1/eps times before destruction and leak out within tau ~ 1/sqrt(eps) of the surface, so J and S fall below B there.",
           "S = eps B + (1-eps) J;  Eddington: J = B [1 - exp(-sqrt(3 eps) tau)/(1 + 2 sqrt(eps/3))];  exact: S(0) = sqrt(eps) B")
header(ws, "A5", "Inputs")
eps = inp(ws, 6, "Destruction probability eps", 0.01, "", "1 = pure absorption (LTE). Try 0.1, 0.01, 0.001.")
header(ws, "A8", "Results")
k3 = f"SQRT(3*{eps})"
den = f"(1+2*SQRT({eps}/3))"
out(ws, 9, "Exact S(0)/B = sqrt(eps)", f"=SQRT({eps})", "", "0.0000")
out(ws, 10, "Eddington S(0)/B", f"={eps}+(1-{eps})*(1-1/{den})", "", "0.0000")
out(ws, 11, "Eddington error at the surface", f"=B10/B9-1", "", "0.0%")
out(ws, 12, "Thermalisation depth 1/sqrt(eps)", f"=1/SQRT({eps})", "", "0.0")
out(ws, 13, "Eddington thermalisation length 1/sqrt(3 eps)", f"=1/{k3}", "", "0.0")
taus = np.logspace(-3, 3.5, 53)
rows = [[float(x), f"={eps}+(1-{eps})*(1-EXP(-{k3}*E{17 + i})/{den})", f"=1-EXP(-{k3}*E{17 + i})/{den}"] for i, x in enumerate(taus)]
r1, r2 = table(ws, 5, 16, ["tau", "S/B (Eddington)", "J/B (Eddington)"], rows, ["0.000E+00", "0.0000", "0.0000"])
chart(ws, "I5", "Source function in an isothermal scattering atmosphere", "optical depth tau (log)", "S/B", 5, [6, 7], r1, r2,
      ["S/B", "J/B"], xlog=True, ymin=0, ymax=1)
ws["A16"] = "Exact numerical S(0)/B (Lambda matrix) for comparison: see sheet Data, columns I-K."
ws["A16"].font = f_note
header(ws, "A19", "Rosseland mean calculator (3 frequency bands)")
ws["A20"] = "Band weights must add to 1 (fraction of dB/dT in each band)."
ws["A20"].font = f_note
w1 = inp(ws, 21, "Weight band 1", 0.5)
w2 = inp(ws, 22, "Weight band 2", 0.3)
w3 = inp(ws, 23, "Weight band 3", 0.2)
k1 = inp(ws, 24, "Opacity band 1", 1, "m^2/kg")
k2 = inp(ws, 25, "Opacity band 2", 10, "m^2/kg")
k3c = inp(ws, 26, "Opacity band 3", 1000, "m^2/kg")
out(ws, 27, "Rosseland mean (harmonic)", f"=({w1}+{w2}+{w3})/({w1}/{k1}+{w2}/{k2}+{w3}/{k3c})", "m^2/kg", "0.000")
out(ws, 28, "Planck-type mean (arithmetic)", f"=({w1}*{k1}+{w2}*{k2}+{w3}*{k3c})/({w1}+{w2}+{w3})", "m^2/kg", "0.000")

# ------------------------------------------------------------------ L9
ws = sheet("SA4 Gray atmosphere", "Lecture SA4 - The complete gray LTE atmosphere",
           "Exact solution (Hopf q, Chandrasekhar H) vs Eddington. The model scales with T_eff; limb darkening and spectrum are integrated numerically.",
           "T^4 = 3/4 T_eff^4 [tau + q(tau)];  I_lambda(0,mu) = Int B_lambda(T(tau)) exp(-tau/mu) dtau/mu;  F_lambda = 2 pi Int B_lambda E2 dtau;  P = g tau / kappa")
header(ws, "A5", "Inputs")
Te = inp(ws, 6, "Effective temperature T_eff", 5772, "K", "Sun 5772 K. The gray solution is self-similar in T/T_eff.")
lamL = inp(ws, 7, "Wavelength for limb darkening", 500, "nm", "Compare 300 nm (strong darkening) with 1500 nm.")
lg = inp(ws, 8, "Surface gravity log g (cgs)", 4.44, "", "Sun 4.44; giants ~2; white dwarfs ~8.")
kap = inp(ws, 9, "Mean opacity kappa", 0.03, "m^2/kg", "Solar photosphere ~0.03 m^2/kg.")
header(ws, "A11", "Results")
out(ws, 12, "Surface T(0), exact", f"={Te}*(0.75*INDEX({Q_R},1))^0.25", "K", "0")
out(ws, 13, "Surface T(0), Eddington", f"={Te}*0.5^0.25", "K", "0")
out(ws, 14, "T at tau = 1 (Eddington)", f"={Te}*(0.75*(1+2/3))^0.25", "K", "0")
P23 = out(ws, 15, "Pressure at tau = 2/3: P = g tau/kappa", f"=10^{lg}/100*(2/3)/{kap}", "Pa", "0.000E+00")
out(ws, 16, "Density at tau = 2/3 (mu_mol = 1.3)", f"={P23}*1.3*mH/(kB*{Te})", "kg m^-3")
# depth table: tau, q, E2, trapezoid weight, T exact, T Eddington, B_lambda(T)
rows = []
r0 = 31
for i in range(NH):
    r = r0 + i
    wexpr = (f"=(E{r + 1}-E{r})/2" if i == 0 else f"=(E{r}-E{r - 1})/2" if i == NH - 1 else f"=(E{r + 1}-E{r - 1})/2")
    rows.append([f"=INDEX({TAU_R},{i + 1})", f"=INDEX({Q_R},{i + 1})", f"=INDEX({E2_R},{i + 1})", wexpr,
                 f"={Te}*(0.75*(E{r}+F{r}))^0.25", f"={Te}*(0.75*(E{r}+2/3))^0.25",
                 f"=2*hP*cL^2/({lamL}*1E-9)^5/(EXP(hP*cL/({lamL}*1E-9*kB*I{r}))-1)"])
d1, d2 = table(ws, 5, r0 - 1, ["tau", "Hopf q", "E2(tau)", "trapezoid weight", "T exact [K]", "T Eddington [K]", "B_lambda(T) at chosen lambda"],
               rows, ["0.000E+00", "0.0000", "0.0000E+00", "0.000E+00", "0", "0", "0.000E+00"])
TAUc, Wc, Bc, Tc, E2c = (f"$E${d1}:$E${d2}", f"$H${d1}:$H${d2}", f"$K${d1}:$K${d2}", f"$I${d1}:$I${d2}", f"$G${d1}:$G${d2}")
out(ws, 17, "Energy check: Int F_lambda dlambda / sigma T_eff^4", "=SUMPRODUCT(T31:T150,U31:U150)/(sigSB*B6^4)", "(should be ~1)", "0.000")
ws["A17"].font = f_norm
# limb darkening table (mu) at chosen wavelength
rows = []
mus = [0.02] + [round(x, 2) for x in np.arange(0.05, 1.0001, 0.05)]
for i, m in enumerate(mus):
    r = 20 + i
    rows.append([m, f"=SUMPRODUCT({Bc},EXP(-{TAUc}/M{r})/M{r},{Wc})", f"=N{r}/N{20 + len(mus) - 1}", f"=(2+3*M{r})/5",
                 f"=INDEX(Data!$G${r1m}:$G${r2m},MATCH(ROUND(M{r}*20,0)/20,Data!$E${r1m}:$E${r2m},1))"])
l1, l2 = table(ws, 13, 19, ["mu", "I_lambda(0,mu)", "I/I(1) at lambda", "Eddington (2+3mu)/5", "bolometric exact"], rows,
               ["0.00", "0.000E+00", "0.0000", "0.0000", "0.0000"])
# spectrum table
rows = []
lamS = np.round(np.logspace(np.log10(20), np.log10(100000), 120), 3)
for i, l in enumerate(lamS):
    r = 31 + i
    wl = (f"=(S{r + 1}-S{r})/2*1E-9" if i == 0 else f"=(S{r}-S{r - 1})/2*1E-9" if i == len(lamS) - 1 else f"=(S{r + 1}-S{r - 1})/2*1E-9")
    rows.append([float(l), f"=2*PI()*SUMPRODUCT(2*hP*cL^2/(S{r}*1E-9)^5/(EXP(hP*cL/(S{r}*1E-9*kB*{Tc}))-1),{E2c},{Wc})",
                 wl, f"=PI()*2*hP*cL^2/(S{r}*1E-9)^5/(EXP(hP*cL/(S{r}*1E-9*kB*{Te}))-1)"])
s1, s2 = table(ws, 19, 30, ["lambda [nm]", "F_lambda gray LTE", "d lambda weight [m]", "pi B_lambda(T_eff)"], rows,
               ["0.0", "0.000E+00", "0.000E+00", "0.000E+00"])
for col, wdt in zip("MNOPQRSTUV", [7, 14, 12, 14, 12, 3, 12, 14, 14, 14]):
    ws.column_dimensions[col].width = wdt
ws["S29"] = "Emergent spectrum (20 nm - 100 um, used for the energy check)"
ws["S29"].font = f_note
chart(ws, "A102", "Temperature structure: exact vs Eddington", "optical depth tau (log)", "T [K]", 5, [9, 10], d1 + 1, d2,
      ["exact (Hopf)", "Eddington"], xlog=True)
chart(ws, "J102", "Limb darkening at the chosen wavelength", "mu", "I(mu)/I(1)", 13, [15, 16, 17], l1, l2,
      ["monochromatic", "Eddington (bolometric)", "exact (bolometric)"], xmin=0, xmax=1, ymin=0, ymax=1)
chart(ws, "J120", "Emergent flux vs black body at T_eff (log-log)", "wavelength [nm]", "F_lambda [W m^-3]", 19, [20, 22], s1 + 3, s2 - 6,
      ["gray LTE atmosphere", "pi B_lambda(T_eff)"], xlog=True, ylog=True)

# ------------------------------------------------------------------ L10
ws = sheet("SA4 Lines", "Lecture SA4 - How does a spectral line grow? Milne-Eddington profiles and the curve of growth",
           "In LTE a line is dark because we see higher, cooler layers at its frequencies. Saturation: the core cannot drop below 1/(1+2beta/3).",
           "F/F_c = [1 + 2beta/(3(1+eta))] / [1 + 2beta/3],  eta = eta0 H(a,v),  W = Int (1 - F/F_c) dv")
header(ws, "A5", "Inputs")
beta = inp(ws, 6, "Temperature gradient beta", 1.5, "", "B = B0 (1 + beta tau). beta = 0: isothermal, no line.")
eta0 = inp(ws, 7, "Line strength eta0 (proportional to N f)", 10, "", "Try 0.1 (weak), 10 (saturated), 1e5 (damping wings).")
ad = inp(ws, 8, "Damping parameter a", 0.01, "", "Voigt damping. Set 0 for a pure Gaussian.")
Tg = inp(ws, 9, "Gas temperature", 5772, "K")
mass = inp(ws, 10, "Atomic mass", 55.8, "u", "Fe 55.8, Ca 40.1, Na 23.0, H 1.008")
xi = inp(ws, 11, "Microturbulence xi", 1, "km/s")
lam0 = inp(ws, 12, "Line wavelength lambda0", 500, "nm")
ws["A13"] = ("Voigt function approximated as H(a,v) ~ exp(-v^2) + a/(sqrt(pi)(1+v^2)), which is exact in the core "
             "and the far wings for a << 1 (the web app uses the exact Voigt function).")
ws["A13"].font = f_note
header(ws, "A15", "Results")
Cc = f"((2*{beta}/3)/(1+2*{beta}/3))"
out(ws, 16, "Central residual flux F/F_c", f"=(1+2*{beta}/(3*(1+{eta0}*(1+{ad}/SQRT(PI())))))/(1+2*{beta}/3)", "", "0.0000")
out(ws, 17, "Maximum possible depth", f"=1-1/(1+2*{beta}/3)", "", "0.0000")
vD = out(ws, 18, "Doppler velocity sqrt(2kT/m + xi^2)", f"=SQRT(2*kB*{Tg}/({mass}*amu)+({xi}*1000)^2)/1000", "km/s", "0.000")
dl = out(ws, 19, "Doppler width (lambda0 v/c)", f"={lam0}*1E-9*{vD}*1000/cL*1E13", "mA", "0.00")
# v-grid for integration (sinh stretched) with weights, in hidden-ish columns P:Q
vg = [float(2 * np.sinh(x)) for x in np.linspace(0, np.arcsinh(500 / 2), 160)]
rows = []
for i, v in enumerate(vg):
    r = 41 + i
    wexpr = (f"=(P{r + 1}-P{r})/2" if i == 0 else f"=(P{r}-P{r - 1})/2" if i == len(vg) - 1 else f"=(P{r + 1}-P{r - 1})/2")
    rows.append([v, wexpr])
v1, v2 = table(ws, 16, 40, ["v (integration grid)", "weight"], rows, ["0.000", "0.0000"])
VG, VW = f"$P${v1}:$P${v2}", f"$Q${v1}:$Q${v2}"
Hexpr = lambda e: f"(EXP(0-{VG}^2)+{ad}/(SQRT(PI())*(1+{VG}^2)))"
Wf = (f"=2*SUMPRODUCT(1-(1+2*{beta}/(3*(1+{eta0}*{Hexpr(0)})))/(1+2*{beta}/3),{VW})")
W = out(ws, 20, "Equivalent width W / Delta nu_D", Wf, "", "0.000")
out(ws, 21, "Equivalent width W_lambda", f"={W}*{dl}", "mA", "0.0")
out(ws, 22, "Weak-line limit sqrt(pi) C eta0", f"=SQRT(PI())*{Cc}*{eta0}", "", "0.000")
# profile table
rows = []
for i in range(81):
    r = 26 + i
    rows.append([f"=-8+{i}*0.2", f"=(1+2*{beta}/(3*(1+{eta0}*(EXP(0-E{r}^2)+{ad}/(SQRT(PI())*(1+E{r}^2))))))/(1+2*{beta}/3)"])
p1, p2 = table(ws, 5, 25, ["v = (nu-nu0)/Delta nu_D", "F/F_c"], rows, ["0.0", "0.0000"])
# curve of growth table (eta0 grid), formula per row
rows = []
etas = np.logspace(-2, 6, 33)
for i, e in enumerate(etas):
    r = 26 + i
    rows.append([float(e), f"=2*SUMPRODUCT(1-(1+2*{beta}/(3*(1+H{r}*{Hexpr(0)})))/(1+2*{beta}/3),{VW})",
                 f"=2*SUMPRODUCT(1-(1+2*{beta}/(3*(1+H{r}*EXP(0-{VG}^2))))/(1+2*{beta}/3),{VW})", f"=SQRT(PI())*{Cc}*H{r}"])
g1, g2 = table(ws, 8, 25, ["eta0", "W/Dnu_D (Voigt approx.)", "W/Dnu_D (Gaussian)", "linear limit"], rows,
               ["0.00E+00", "0.0000", "0.0000", "0.0000"])
for col, wdt in zip("HIJKP", [11, 14, 14, 12, 12]):
    ws.column_dimensions[col].width = wdt
chart(ws, "S5", "Line profile (LTE, Milne-Eddington)", "v = (nu - nu0)/Delta nu_D", "F/F_c", 5, [6], p1, p2, ["F/F_c"], xmin=-8, xmax=8)
chart(ws, "S23", "Curve of growth (log-log)", "eta0 (proportional to number of absorbers)", "W / Delta nu_D", 8, [9, 10, 11], g1, g2,
      ["with damping", "Gaussian only", "linear limit"], xlog=True, ylog=True, ymin=0.001, ymax=1000)

# order: put Start first, Constants and Data last
order = ["Start here"] + [s for s, _ in SHEETS[:10]] + ["Constants", "Data"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print("wrote", OUT)
